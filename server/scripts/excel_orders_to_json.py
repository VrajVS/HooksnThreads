"""Convert the studio's order spreadsheet (Crochet Orders.xlsx) into orders JSON.

    python scripts/excel_orders_to_json.py "Crochet Orders.xlsx" orders.json

Needs openpyxl (not a server dependency: run it on a PC). The JSON is then
loaded with scripts/load_orders_json.py. A report of every row that needed
interpreting is printed; check it before loading.

Sheet layout ("Orders 2024/2025/2026"): month header rows ("JANUARY"), repeated
column-header rows, "Total" rows, and order rows. A row with an invoice number
(or exhibition name) or a client starts an order; following rows with neither
are more items of that order.
"""

import datetime as dt
import json
import re
import sys
from collections import defaultdict

import openpyxl

MONTHS = {m: i for i, m in enumerate(
    ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST",
     "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"], start=1)}
SHEETS = {"Orders 2024": 2024, "Orders 2025": 2025, "Orders 2026": 2026}
CHARGE_WORDS = {"wrapping": "Wrapping", "courier": "Courier", "shipping": "Shipping", "delivery": "Delivery"}
SOURCE_TAG = "Imported from Crochet Orders.xlsx"

report: dict[str, list[str]] = defaultdict(list)


def note(kind: str, msg: str) -> None:
    report[kind].append(msg)


def clean(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def yes(v) -> bool:
    return clean(v) is not None and clean(v).lower() in ("yes", "yea", "y")


def num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    return None


def parse_qty(v, where) -> tuple[int, str | None]:
    """Returns (quantity, detail to append to the title)."""
    if isinstance(v, (int, float)):
        return max(1, int(v)), None
    s = clean(v)
    if not s:
        return 1, None
    m = re.match(r"^(\d+)\s*(\(.*\))?$", s)
    if m:
        return int(m.group(1)), m.group(2)
    note("quantity", f"{where}: quantity '{s}' -> 1, kept in the description")
    return 1, f"({s})"


def parse_price(v, where):
    """Returns (unit_price, charges[(label, amount|None)], title_note, all_unlabeled_parts)."""
    if isinstance(v, (int, float)):
        return float(v), [], None, ([], False)
    s = clean(v)
    if not s:
        return None, [], None, ([], False)
    is_list = "," in s and "+" not in s
    parts = [p.strip() for p in re.split(r"[+,]", s) if p.strip()]
    item_parts, charges, notes = [], [], []
    for p in parts:
        m = re.match(r"^(\d+(?:\.\d+)?)\s*(?:\((.*)\))?\s*(.*)$", p)
        if m:
            amount, paren, rest = float(m.group(1)), (m.group(2) or "").strip(), m.group(3).strip()
            label_text = (paren + " " + rest).strip().lower()
            charge = next((lbl for w, lbl in CHARGE_WORDS.items() if w in label_text), None)
            if charge:
                charges.append((charge, amount))
            else:
                item_parts.append(amount)
                if paren:
                    notes.append(paren)
        else:
            word = p.lower()
            charge = next((lbl for w, lbl in CHARGE_WORDS.items() if w in word), None)
            if charge:
                charges.append((charge, None))  # amount taken from the row total
            else:
                notes.append(p)
    unit = sum(item_parts) if item_parts else None
    note("price", f"{where}: price '{s}' -> item {unit:g}" + (
        "".join(f", {c} {a:g}" if a is not None else f", {c} (from total)" for c, a in charges)) +
        (f", note '{'; '.join(notes)}'" if notes else "") if unit is not None else f"{where}: price '{s}' unreadable")
    return unit, charges, "; ".join(notes) or None, (item_parts, is_list)


def parse_delivery(v, year: int, where) -> str | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        # Typed as d/m/yyyy, but Excel stored it as m/d: swap back.
        try:
            fixed = dt.date(v.year, v.day, v.month)
        except ValueError:
            fixed = v.date()
        return fixed.isoformat()
    s = clean(v)
    s_end = s.split("-")[-1].strip()  # "20/3 - 22/3" -> last day of the range
    m = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?$", s_end)
    if not m:
        note("dates", f"{where}: delivery date '{s}' not understood, left blank")
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = year if y is None else (int(y) + 2000 if len(y) == 2 else int(y))
    if y != year and abs(y - year) > 1:
        note("dates", f"{where}: delivery year {y} looks like a typo, using {year}")
        y = year
    if y != year and s_end != s:
        y = year
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        note("dates", f"{where}: delivery date '{s}' is not a real date, left blank")
        return None


def mode_of(v) -> str | None:
    s = (clean(v) or "").lower()
    if "online" in s and "cash" in s:
        return "online_cash"
    if "online" in s:
        return "online"
    if "cash" in s:
        return "cash"
    return None


def read_sheet(ws, year):
    orders, cur, month = [], None, None
    for r in range(2, ws.max_row + 1):
        a, item, client, qty, prep, deliv, price, total, recv, mode, ddate = (
            ws.cell(r, c).value for c in range(1, 12))
        where = f"{ws.title} row {r}"
        if all(x is None for x in (a, item, client, qty, price, total)):
            continue
        if isinstance(a, str) and clean(a).upper() in MONTHS and not item:
            month = MONTHS[clean(a).upper()]
            cur = None
            continue
        if clean(a) == "Invoice #" or clean(item) == "Item":
            continue
        if clean(price) and clean(price).lower() in ("total", "grand total"):
            continue
        if not clean(item):
            if clean(a) or total not in (None, 0):
                note("skipped", f"{where}: '{clean(a) or total}' (no item) skipped")
            continue
        if month is None:
            note("skipped", f"{where}: row before any month header skipped")
            continue

        label = clean(a) if isinstance(a, str) else None
        starts = a is not None or clean(client) is not None
        if not starts and cur is None:
            if price is None and total is None:
                note("skipped", f"{where}: note '{clean(item)}' (no order, no price) skipped")
                continue
            starts = True
        if starts:
            exhibition = bool(label and "exhibition" in label.lower() or label in ("Himalaya Mall",))
            cur = {
                "sheet": ws.title, "row": r, "year": year, "month": month,
                "invoice": str(int(a)) if isinstance(a, (int, float)) else None,
                "label": label, "exhibition": exhibition,
                # A named exhibition block is one stall's sales; "IDK" means unknown.
                "customer_name": (label if exhibition and label.lower() != "exhibition" else None)
                or (None if (clean(client) or "").upper() == "IDK" else clean(client))
                or ("Exhibition walk-in" if exhibition else label) or "Walk-in customer",
                "items": [], "charges": [], "rows": [],
                "first_total": num(total), "first_recv": recv, "first_mode": mode,
                "delivery_date": None, "notes": [],
            }
            orders.append(cur)
        q, qty_note = parse_qty(qty, where)
        unit, charges, price_note, (parts, is_list) = parse_price(price, where)
        line_total = num(total)
        title = clean(item)
        if is_list and len(parts) == q:
            # One price per piece: a line per distinct price.
            for extra in (qty_note,):
                if extra:
                    title = f"{title} {extra}"
            counts = defaultdict(int)
            for part in parts:
                counts[part] += 1
            for each, n in counts.items():
                cur["items"].append({
                    "title": title, "quantity": n, "unit_price": each, "line_total": None,
                    "prepared": yes(prep), "delivered": yes(deliv), "received": recv,
                    "mode": mode_of(mode), "advance": None, "sheet_row": r, "plain_price": False,
                })
            note("price", f"{where}: '{title}' split into " + ", ".join(f"{n} x {e:g}" for e, n in counts.items()))
            dd = parse_delivery(ddate, year, where)
            if dd and not cur["delivery_date"]:
                cur["delivery_date"] = dd
            continue
        if unit is not None and isinstance(price, str) and len(parts) + len(charges) > 1 and q > 1:
            note("price", f"{where}: '{title}' price '{clean(price)}' is for the whole line -> {unit / q:g} each x {q}")
            unit = unit / q
        for extra in (qty_note, f"({price_note})" if price_note else None):
            if extra:
                title = f"{title} {extra}"
        advance = None
        if isinstance(prep, str) and "advance" in prep.lower():
            m = re.search(r"(\d+)", prep)
            advance = float(m.group(1)) if m else None
            note("payments", f"{where}: '{clean(prep)}' -> advance payment {advance:g}, not yet prepared")
        cur["items"].append({
            "title": title, "quantity": q, "unit_price": unit, "line_total": line_total,
            "prepared": yes(prep), "delivered": yes(deliv),
            "received": recv, "mode": mode_of(mode), "advance": advance, "sheet_row": r,
            "plain_price": isinstance(price, (int, float)) and not charges,
        })
        cur["charges"].extend(charges)
        dd = parse_delivery(ddate, year, where)
        if dd and not cur["delivery_date"]:
            cur["delivery_date"] = dd
    return orders


def finalise(o):
    where = f"{o['sheet']} row {o['row']}"
    items = o["items"]
    # The Total column is either per line (every item row filled in, as at
    # exhibitions) or one figure for the whole order on its first row.
    per_line = len(items) > 1 and any(it["line_total"] is not None for it in items[1:])
    for it in items:
        lt = it["line_total"] if (per_line or len(items) == 1) else None
        if (it["plain_price"] and lt is not None and it["unit_price"] is not None
                and round(it["unit_price"] * it["quantity"]) != round(lt) and lt % it["quantity"] == 0):
            note("price", f"{where}: '{it['title']}' {it['quantity']} x {it['unit_price']:g} != row total {lt:g} -> {lt / it['quantity']:g} each")
            it["unit_price"] = lt / it["quantity"]
    for it in items:
        if it["unit_price"] is None:
            if it["line_total"] is not None:
                it["unit_price"] = it["line_total"] / it["quantity"]
                note("price", f"{where}: '{it['title']}' price missing, used row total")
            else:
                it["unit_price"] = 0
                note("price", f"{where}: '{it['title']}' has no price or total, set to 0")
        it["unit_price"] = round(it["unit_price"])

    subtotal = sum(it["unit_price"] * it["quantity"] for it in items)
    known = sum(a for _, a in o["charges"] if a is not None)
    if per_line:
        sheet_total = sum(
            it["line_total"] if it["line_total"] is not None else it["unit_price"] * it["quantity"]
            for it in items
        )
    else:
        sheet_total = o["first_total"]
    charges = [(lbl, a) for lbl, a in o["charges"] if a is not None]
    missing = [lbl for lbl, a in o["charges"] if a is None]
    if sheet_total is not None:
        diff = round(sheet_total - subtotal - known)
        if missing:
            charges.append((missing[0], diff))
            diff = 0
        if diff:
            label = "Cashback" if diff < 0 and any("cashback" in it["title"].lower() for it in items) else "Adjustment to match sheet total"
            charges.append((label, diff))
            note("totals", f"{where} ({o['customer_name']}): items + charges = {subtotal + known:g}, sheet total {sheet_total:g} -> added '{label}' {diff:+g}")
    total = subtotal + sum(a for _, a in charges)
    if total < 0:
        note("totals", f"{where}: negative total, charges dropped")
        charges, total = [], subtotal

    # Payments
    payments, payment_mode = [], None
    if o["exhibition"]:
        by_mode = defaultdict(float)
        fallback = mode_of(o["first_mode"]) or "cash"
        for it in items:
            by_mode[it["mode"] or fallback] += it["unit_price"] * it["quantity"]
        for m, amt in by_mode.items():
            if amt > 0:
                payments.append({"mode": m, "amount": amt})
        gap = total - sum(p["amount"] for p in payments)
        if payments and gap:
            max(payments, key=lambda p: p["amount"])["amount"] += gap
        payment_mode = "online_cash" if len(by_mode) > 1 or fallback == "online_cash" else fallback
        for it in items:
            it["prepared"] = it["delivered"] = True
    else:
        payment_mode = mode_of(o["first_mode"])
        recv = o["first_recv"]
        if isinstance(recv, (int, float)) and recv > 0:
            payments.append({"mode": payment_mode or "cash", "amount": float(recv), "note": "Part payment (from sheet)"})
            note("payments", f"{where}: amount received {recv:g} of {total:g}")
        elif yes(recv) and total > 0:
            payments.append({"mode": payment_mode or "cash", "amount": float(total)})
        for it in items:
            if it["advance"]:
                payments.append({"mode": payment_mode or "cash", "amount": it["advance"], "note": "Advance"})
    split_payments = []
    for p in payments:
        if p["mode"] == "online_cash":
            note("payments", f"{where} ({o['customer_name']}): {p['amount']:g} paid 'Online + Cash', split unknown -> recorded as Online")
            p = {**p, "mode": "online", "note": "Online + Cash (split not recorded in the sheet)"}
        split_payments.append({**p, "amount": round(p["amount"])})

    delivered = all(it["delivered"] for it in items)
    paid = sum(p["amount"] for p in split_payments)
    prepared_all = all(it["prepared"] for it in items)
    status = "completed" if delivered and prepared_all and paid >= total else "confirmed"
    notes = [f"{SOURCE_TAG} ({o['sheet']}, row {o['row']})."]
    if o["exhibition"]:
        notes.append(f"Exhibition sales: {o['label']}.")
    elif o["label"]:
        notes.append(f"Sheet label: {o['label']}.")
    return {
        "invoice_number": o["invoice"],
        "order_date": dt.date(o["year"], o["month"], 1).isoformat(),
        "customer_name": o["customer_name"],
        "status": status,
        "delivered": delivered,
        "delivery_date": o["delivery_date"],
        "payment_mode": payment_mode,
        "notes": " ".join(notes),
        "items": [{"title": it["title"], "quantity": it["quantity"], "unit_price": it["unit_price"],
                   "prepared": it["prepared"]} for it in items],
        "charges": [{"label": l, "amount": a} for l, a in charges if a],
        "payments": split_payments,
        "_where": where,
        "_exhibition": o["exhibition"],
        "_year": o["year"],
        "_month": o["month"],
    }


def assign_invoice_numbers(orders):
    used = set()
    for o in orders:
        # Early 2024 used a plain running count (1, 2, 3...); give those the
        # YYYYMM prefix the sheet switched to later, keeping the count.
        if o["invoice_number"] and len(o["invoice_number"]) < 5:
            o["invoice_number"] = f"{o['_year']}{o['_month']:02d}{int(o['invoice_number']):02d}"
        if o["invoice_number"]:
            if o["invoice_number"] in used:
                note("invoices", f"{o['_where']}: invoice {o['invoice_number']} used twice, second one gets a suffix")
                o["invoice_number"] += "-2"
            used.add(o["invoice_number"])
    seq = defaultdict(int)
    ex_seq = defaultdict(int)
    for o in orders:
        if o["invoice_number"]:
            continue
        ym = f"{o['_year']}{o['_month']:02d}"
        if o["_exhibition"]:
            ex_seq[ym] += 1
            o["invoice_number"] = f"EX{ym}{ex_seq[ym]:02d}"
        else:
            while True:
                seq[ym] += 1
                cand = f"{ym}{seq[ym]:02d}"
                if cand not in used:
                    break
            o["invoice_number"] = cand
            if o["_year"] != 2024:
                note("invoices", f"{o['_where']} ({o['customer_name']}): no invoice number in sheet -> {cand}")
        used.add(o["invoice_number"])


def attach_pending(wb, orders):
    ws = wb["Pending orders"]
    client = None
    for r in range(2, ws.max_row + 1):
        sr, c, order_text, todo, qty = (ws.cell(r, k).value for k in range(1, 6))
        if clean(c):
            client = clean(c)
        text = clean(order_text)
        if not text or not client:
            continue
        first = client.split("(")[0].strip().lower()
        words = {w for w in re.findall(r"[a-z]+", text.lower()) if len(w) > 2}
        best, best_score = None, 0
        for o in orders:
            if not o["customer_name"].lower().startswith(first.split()[0]):
                continue
            for it in o["items"]:
                iw = set(re.findall(r"[a-z]+", it["title"].lower()))
                overlap = len(words & iw)
                if not overlap:
                    continue
                score = overlap + (2 if not it["prepared"] else 0)
                if score > best_score:
                    best, best_score = (o, it), score
        if best and best_score >= 2:
            o, it = best
            it["prepared"] = False
            if clean(todo):
                it["work_note"] = clean(todo)
            o["status"] = "confirmed"
            note("pending", f"Pending row {r} ({client}: {text}) -> {o['invoice_number']} '{it['title']}'" +
                 (f", note '{clean(todo)}'" if clean(todo) else ""))
        else:
            note("pending-unmatched", f"Pending row {r} ({client}: {text}{', ' + clean(todo) if clean(todo) else ''}) - no matching order found")


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)
    raw = []
    for name, year in SHEETS.items():
        raw.extend(read_sheet(wb[name], year))
    orders = [finalise(o) for o in raw]
    assign_invoice_numbers(orders)
    attach_pending(wb, orders)
    for o in orders:
        for k in [k for k in o if k.startswith("_")]:
            del o[k]
    with open(out, "w", encoding="utf-8") as f:
        json.dump(orders, f, ensure_ascii=False, indent=1)

    total_items = sum(len(o["items"]) for o in orders)
    revenue = sum(sum(i["unit_price"] * i["quantity"] for i in o["items"]) + sum(c["amount"] for c in o["charges"]) for o in orders)
    print(f"{len(orders)} orders, {total_items} item lines, total value {revenue:,.0f}")
    by_year = defaultdict(lambda: [0, 0])
    for o in orders:
        y = o["order_date"][:4]
        by_year[y][0] += 1
        by_year[y][1] += sum(i["unit_price"] * i["quantity"] for i in o["items"]) + sum(c["amount"] for c in o["charges"])
    for y, (n, v) in sorted(by_year.items()):
        print(f"  {y}: {n} orders, value {v:,.0f}")
    for kind, msgs in report.items():
        print(f"\n## {kind} ({len(msgs)})")
        for m in msgs:
            print("  " + m)


if __name__ == "__main__":
    main()
